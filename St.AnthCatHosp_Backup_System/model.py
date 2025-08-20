def create_tables(conn):
    cursor = conn.cursor()

    # Enable foreign key support (necessary in SQLite)
    cursor.execute("PRAGMA foreign_keys = ON;")

    # Create departments first (used as FK in users and files)
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS departments (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT UNIQUE NOT NULL
    );
    """)

    # Users table
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS users (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT UNIQUE NOT NULL,
        password TEXT NOT NULL,
        role TEXT NOT NULL CHECK(role IN ('chief_admin', 'department_admin', 'user')),
        department_id INTEGER,
        FOREIGN KEY(department_id) REFERENCES departments(id)
    );
    """)

    # Files table
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS files (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        filename TEXT NOT NULL,
        content BLOB NOT NULL,
        uploaded_by INTEGER,
        department_id INTEGER,
        timestamp DATETIME DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY(uploaded_by) REFERENCES users(id),
        FOREIGN KEY(department_id) REFERENCES departments(id)
    );
    """)

    # Audit log table
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS audit_log (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        user_id INTEGER,
        file_id INTEGER,
        action TEXT NOT NULL,
        timestamp DATETIME DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY(user_id) REFERENCES users(id),
        FOREIGN KEY(file_id) REFERENCES files(id)
    );
    """)

    conn.commit()

import streamlit as st
from db import get_connection
import mimetypes
import datetime
import io
import base64
import pandas as pd

try:
    import docx
    import openpyxl
except ImportError:
    st.warning("Please install 'python-docx' and 'openpyxl' for .docx/.xlsx previews.")

ITEMS_PER_PAGE = 5

def show(user):
    st.title("📁 Dashboard - Uploaded Files")

    conn = get_connection()
    cursor = conn.cursor()

    # Department filter (chief_admin only)
    selected_department = None
    if user["role"] == "chief_admin":
        cursor.execute("SELECT id, name FROM departments")
        departments = cursor.fetchall()
        dept_dict = {name: id for id, name in departments}
        dept_names = list(dept_dict.keys())
        selected = st.selectbox("📌 Filter by Department", ["All"] + dept_names)
        if selected != "All":
            selected_department = dept_dict[selected]

    # Date range filter
    col1, col2 = st.columns(2)
    with col1:
        start_date = st.date_input("📅 Start date", datetime.date(2000, 1, 1))
    with col2:
        end_date = st.date_input("📅 End date", datetime.date.today())

    # Search box
    search_query = st.text_input("🔍 Search by file name")

    # Fetch files
    query = "SELECT id, filename, timestamp, department_id FROM files WHERE date(timestamp) BETWEEN ? AND ?"
    params = [start_date, end_date]

    if user["role"] != "chief_admin":
        query += " AND department_id = ?"
        params.append(user["department_id"])
    elif selected_department:
        query += " AND department_id = ?"
        params.append(selected_department)

    cursor.execute(query + " ORDER BY timestamp DESC", tuple(params))
    files = cursor.fetchall()
    conn.close()

    # Filter by filename
    if search_query:
        files = [f for f in files if search_query.lower() in f[1].lower()]

    if not files:
        st.warning("No matching files found.")
        return

    # Pagination
    total_pages = (len(files) - 1) // ITEMS_PER_PAGE + 1
    page = st.number_input("Page", min_value=1, max_value=total_pages, value=1)
    start_idx = (page - 1) * ITEMS_PER_PAGE
    end_idx = start_idx + ITEMS_PER_PAGE
    current_files = files[start_idx:end_idx]

    for file_id, filename, timestamp, dept_id in current_files:
        st.markdown(f"### 📄 {filename}")
        st.caption(f"🕒 Uploaded on: {timestamp}")

        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT content FROM files WHERE id = ?", (file_id,))
        result = cursor.fetchone()
        conn.close()

        if not result:
            st.error("⚠️ File content not found.")
            continue

        content = result[0]
        mime_type, _ = mimetypes.guess_type(filename)
        mime_type = mime_type or "application/octet-stream"

        # Preview + download based on MIME
        if mime_type.startswith("image/"):
            st.image(content, caption=filename)

        elif mime_type == "application/pdf":
            b64 = base64.b64encode(content).decode("utf-8")
            st.markdown(f'<iframe src="data:application/pdf;base64,{b64}" width="700" height="500"></iframe>', unsafe_allow_html=True)

        elif mime_type.startswith("text/"):
            try:
                st.text(content.decode("utf-8"))
            except:
                st.warning("Encoding issue in text preview.")

        elif filename.endswith(".docx"):
            try:
                doc = docx.Document(io.BytesIO(content))
                st.markdown("#### 📃 Document Preview:")
                for para in doc.paragraphs:
                    st.write(para.text)
            except:
                st.warning("Unable to preview DOCX.")

        elif filename.endswith(".xlsx"):
            try:
                wb = openpyxl.load_workbook(io.BytesIO(content))
                sheet = wb.active
                data = sheet.values
                df = pd.DataFrame(data)
                st.markdown("#### 📊 Excel Preview:")
                st.dataframe(df)
            except:
                st.warning("Unable to preview XLSX.")

        # Download button for all
        st.download_button(
            label="⬇️ Download",
            data=content,
            file_name=filename,
            mime=mime_type,
            key=f"download_{file_id}_{page}"
        )

    st.caption(f"Page {page} of {total_pages}")

