import streamlit as st
from db import get_connection
import mimetypes
import datetime
import io
import base64
import pandas as pd

try:
    import docx
    import openpyxl
except ImportError:
    st.warning("Please install 'python-docx' and 'openpyxl' for .docx/.xlsx previews.")

ITEMS_PER_PAGE = 5

def show(user):
    st.title("📁 Dashboard - Uploaded Files")

    conn = get_connection()
    cursor = conn.cursor()

    # Department filter (chief_admin only)
    selected_department = None
    if user["role"] == "chief_admin":
        cursor.execute("SELECT id, name FROM departments")
        departments = cursor.fetchall()
        dept_dict = {name: id for id, name in departments}
        dept_names = list(dept_dict.keys())
        selected = st.selectbox("📌 Filter by Department", ["All"] + dept_names)
        if selected != "All":
            selected_department = dept_dict[selected]

    # Date range filter
    col1, col2 = st.columns(2)
    with col1:
        start_date = st.date_input("📅 Start date", datetime.date(2000, 1, 1))
    with col2:
        end_date = st.date_input("📅 End date", datetime.date.today())

    # Search box
    search_query = st.text_input("🔍 Search by file name")

    # Fetch files
    query = "SELECT id, filename, timestamp, department_id FROM files WHERE date(timestamp) BETWEEN ? AND ?"
    params = [start_date, end_date]

    if user["role"] != "chief_admin":
        query += " AND department_id = ?"
        params.append(user["department_id"])
    elif selected_department:
        query += " AND department_id = ?"
        params.append(selected_department)

    cursor.execute(query + " ORDER BY timestamp DESC", tuple(params))
    files = cursor.fetchall()
    conn.close()

    # Filter by filename
    if search_query:
        files = [f for f in files if search_query.lower() in f[1].lower()]

    if not files:
        st.warning("No matching files found.")
        return

    # Pagination
    total_pages = (len(files) - 1) // ITEMS_PER_PAGE + 1
    page = st.number_input("Page", min_value=1, max_value=total_pages, value=1)
    start_idx = (page - 1) * ITEMS_PER_PAGE
    end_idx = start_idx + ITEMS_PER_PAGE
    current_files = files[start_idx:end_idx]

    for file_id, filename, timestamp, dept_id in current_files:
        st.markdown(f"### 📄 {filename}")
        st.caption(f"🕒 Uploaded on: {timestamp}")

        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT content FROM files WHERE id = ?", (file_id,))
        result = cursor.fetchone()
        conn.close()

        if not result:
            st.error("⚠️ File content not found.")
            continue

        content = result[0]
        mime_type, _ = mimetypes.guess_type(filename)
        mime_type = mime_type or "application/octet-stream"

        # Preview + download based on MIME
        if mime_type.startswith("image/"):
            st.image(content, caption=filename)

        elif mime_type == "application/pdf":
            b64 = base64.b64encode(content).decode("utf-8")
            st.markdown(f'<iframe src="data:application/pdf;base64,{b64}" width="700" height="500"></iframe>', unsafe_allow_html=True)

        elif mime_type.startswith("text/"):
            try:
                st.text(content.decode("utf-8"))
            except:
                st.warning("Encoding issue in text preview.")

        elif filename.endswith(".docx"):
            try:
                doc = docx.Document(io.BytesIO(content))
                st.markdown("#### 📃 Document Preview:")
                for para in doc.paragraphs:
                    st.write(para.text)
            except:
                st.warning("Unable to preview DOCX.")

        elif filename.endswith(".xlsx"):
            try:
                wb = openpyxl.load_workbook(io.BytesIO(content))
                sheet = wb.active
                data = sheet.values
                df = pd.DataFrame(data)
                st.markdown("#### 📊 Excel Preview:")
                st.dataframe(df)
            except:
                st.warning("Unable to preview XLSX.")

        # Download button for all
        st.download_button(
            label="⬇️ Download",
            data=content,
            file_name=filename,
            mime=mime_type,
            key=f"download_{file_id}_{page}"
        )

    st.caption(f"Page {page} of {total_pages}")


import streamlit as st
from db import init_db
from db import create_default_admin
from auth import login, logout, get_current_user
from views import (
    dashboard,
    manage_users,
    manage_departments,
    file_upload,
    audit_log,
    export_data
)

# Initialize database
init_db()
create_default_admin()

st.set_page_config(page_title="Department File Backup System", layout="wide")

def render_sidebar(user):
    st.sidebar.title(f"Welcome, {user['username']} ({user['role']})")
    menu_options = ["Dashboard", "Upload File"]

    if user["role"] == "chief_admin":
        menu_options += ["Manage Users", "Manage Departments", "Audit Log", "Export Data"]
    elif user["role"] == "department_admin":
        menu_options += ["Manage Users", "Audit Log"]

    menu_options.append("Logout")
    return st.sidebar.radio("Go to", menu_options)

def main():
    user = get_current_user()  # ✅ Get the user dynamically inside main()

    if not user:
        login()  # Show login form
        return

    # Sidebar and navigation
    selected_page = render_sidebar(user)

    if selected_page == "Dashboard":
        dashboard.show(user)
    elif selected_page == "Upload File":
        file_upload.show(user)
    elif selected_page == "Manage Users" and user["role"] == "chief_admin":
        manage_users.show()
    elif selected_page == "Manage Departments" and user["role"] == "chief_admin":
        manage_departments.show()
    elif selected_page == "Audit Log":
        audit_log.show(user)
    elif selected_page == "Export Data":
        export_data.show(user)
    elif selected_page == "Logout":
        logout()
        st.experimental_rerun()

if __name__ == "__main__":
    main()

